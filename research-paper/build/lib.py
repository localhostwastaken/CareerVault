import re
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement, parse_xml

FONT = "Times New Roman"
M = "http://schemas.openxmlformats.org/officeDocument/2006/math"
W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"

def _rfonts(el, name=FONT):
    rpr = el.get_or_add_rPr() if hasattr(el, "get_or_add_rPr") else el
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.insert(0, rf)
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(a), name)
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if rf.get(qn(a)) is not None: del rf.attrib[qn(a)]

def style_font(style, size, bold=False, italic=False, smallcaps=False, name=FONT):
    f = style.font
    f.name = name; f.size = Pt(size); f.bold = bold; f.italic = italic; f.small_caps = smallcaps
    f.color.rgb = RGBColor(0, 0, 0)
    _rfonts(style.element, name)

class Paper:
    def __init__(self):
        self.doc = Document()
        self.refs = {}      # key -> text
        self.order = []     # keys in first-citation order
        self.fig_n = 0; self.tab_n = 0; self.eq_n = 0
        self._setup()

    # ---------- setup ----------
    def _setup(self):
        d = self.doc
        s = d.sections[0]
        s.page_width = Inches(8.5); s.page_height = Inches(11)
        s.top_margin = Inches(0.75); s.bottom_margin = Inches(1.0)
        s.left_margin = Inches(0.625); s.right_margin = Inches(0.625)
        s.header_distance = Inches(0.4); s.footer_distance = Inches(0.5)
        st = d.styles
        n = st["Normal"]; style_font(n, 10)
        pf = n.paragraph_format; pf.space_before = Pt(0); pf.space_after = Pt(0); pf.line_spacing = 1.0
        self._mk("Body Text IEEE", 10, align=WD_ALIGN_PARAGRAPH.JUSTIFY, first=0.2, after=0)
        self._mk("Abstract IEEE", 9, bold=True, align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=6)
        self._mk("Reference IEEE", 8, align=WD_ALIGN_PARAGRAPH.JUSTIFY, left=0.3, first=-0.3, after=2)
        self._mk("Table Text IEEE", 7.5, align=WD_ALIGN_PARAGRAPH.LEFT, after=0)
        self._mk("Author IEEE", 10, align=WD_ALIGN_PARAGRAPH.CENTER)
        self._mk("Equation IEEE", 10, align=WD_ALIGN_PARAGRAPH.LEFT, before=3, after=3)
        self._mk("Bullet IEEE", 10, align=WD_ALIGN_PARAGRAPH.JUSTIFY, left=0.25, first=-0.18, after=1)
        t = st["Title"]; style_font(t, 24)
        t.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        t.paragraph_format.space_after = Pt(8); t.paragraph_format.space_before = Pt(0)
        ppr = t.element.get_or_add_pPr()
        for b in ppr.findall(qn("w:pBdr")): ppr.remove(b)
        h1 = st["Heading 1"]; style_font(h1, 10, smallcaps=True)
        h1.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        h1.paragraph_format.space_before = Pt(12); h1.paragraph_format.space_after = Pt(4)
        h1.paragraph_format.keep_with_next = True
        h2 = st["Heading 2"]; style_font(h2, 10, italic=True)
        h2.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        h2.paragraph_format.space_before = Pt(6); h2.paragraph_format.space_after = Pt(3)
        h2.paragraph_format.keep_with_next = True
        c = st["Caption"]; style_font(c, 8)
        c.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        c.paragraph_format.space_before = Pt(2); c.paragraph_format.space_after = Pt(8)
        # footer page number
        fp = s.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self._field(fp, "PAGE", 8)

    def _mk(self, name, size, bold=False, italic=False, align=None, left=None, first=None, before=0, after=0):
        from docx.enum.style import WD_STYLE_TYPE
        s = self.doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        s.base_style = self.doc.styles["Normal"]
        style_font(s, size, bold, italic)
        pf = s.paragraph_format
        if align is not None: pf.alignment = align
        if left is not None: pf.left_indent = Inches(left)
        if first is not None: pf.first_line_indent = Inches(first)
        pf.space_before = Pt(before); pf.space_after = Pt(after)
        return s

    def _field(self, p, code, size):
        for typ, txt in (("begin", None), (None, code), ("end", None)):
            r = p.add_run(); r.font.size = Pt(size); r.font.name = FONT
            if typ:
                fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), typ); r._r.append(fc)
            else:
                it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = txt; r._r.append(it)

    # ---------- rich text ----------
    def cite(self, m):
        keys = [k.strip().lstrip("@") for k in m.group(1).split(";")]
        nums = []
        for k in keys:
            if k not in self.refs: raise KeyError("unknown ref " + k)
            if k not in self.order: self.order.append(k)
            nums.append(self.order.index(k) + 1)
        return ", ".join(f"[{n}]" for n in nums)

    TOK = re.compile(r"(\*\*.+?\*\*|\*.+?\*|`.+?`|\^\{.+?\}|_\{.+?\})")
    def rich(self, p, text, size=None, italic=False, bold=False):
        text = re.sub(r"\[(@[^\]]+)\]", self.cite, text)
        for part in self.TOK.split(text):
            if not part: continue
            r = None
            if part.startswith("**"): r = p.add_run(part[2:-2]); r.bold = True
            elif part.startswith("*"): r = p.add_run(part[1:-1]); r.italic = True
            elif part.startswith("`"):
                r = p.add_run(part[1:-1]); r.font.name = "Courier New"; _rfonts(r._r, "Courier New")
                r.font.size = Pt((size or 10) - 1.5)
            elif part.startswith("^{"): r = p.add_run(part[2:-1]); r.font.superscript = True
            elif part.startswith("_{"): r = p.add_run(part[2:-1]); r.font.subscript = True
            else: r = p.add_run(part)
            if r.font.name is None or r.font.name != "Courier New":
                r.font.name = FONT; _rfonts(r._r)
            if size and not part.startswith("`"): r.font.size = Pt(size)
            if italic: r.italic = True
            if bold: r.bold = True
        return p

    # ---------- blocks ----------
    def para(self, text, style="Body Text IEEE", first_after_heading=False, **kw):
        p = self.doc.add_paragraph(style=style)
        if first_after_heading: p.paragraph_format.first_line_indent = Inches(0)
        self.rich(p, text, **kw)
        return p

    def h1(self, text): return self.doc.add_paragraph(text, style="Heading 1")
    def h2(self, text): return self.doc.add_paragraph(text, style="Heading 2")

    def bullets(self, items):
        for i, t in enumerate(items):
            p = self.doc.add_paragraph(style="Bullet IEEE")
            self.rich(p, f"({i+1}) " + t if False else "• " + t)

    def figure(self, path, caption, width=3.4):
        self.fig_n += 1
        p = self.doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True; p.paragraph_format.space_before = Pt(4)
        p.add_run().add_picture(path, width=Inches(width))
        c = self.doc.add_paragraph(style="Caption")
        self.rich(c, f"Fig. {self.fig_n}. " + caption, size=8)
        return self.fig_n

    ROMAN = ["", "I", "II", "III", "IV", "V", "VI", "VII", "VIII", "IX", "X", "XI", "XII"]
    def table(self, caption, header, rows, widths, align=None):
        self.tab_n += 1
        cp = self.doc.add_paragraph(style="Caption"); cp.paragraph_format.keep_with_next = True
        cp.paragraph_format.space_after = Pt(3); cp.paragraph_format.space_before = Pt(6)
        r = cp.add_run(f"TABLE {self.ROMAN[self.tab_n]}"); r.font.size = Pt(8); r.font.small_caps = False
        cp.add_run().add_break()
        self.rich(cp, caption.upper() if False else caption, size=8)
        for rr in cp.runs[1:]: rr.font.small_caps = True
        t = self.doc.add_table(rows=1 + len(rows), cols=len(header))
        t.alignment = WD_TABLE_ALIGNMENT.LEFT; t.autofit = False
        tblPr = t._tbl.tblPr
        for e in tblPr.findall(qn("w:tblStyle")): tblPr.remove(e)
        borders = parse_xml(f'<w:tblBorders xmlns:w="{W}"><w:top w:val="single" w:sz="8" w:space="0" w:color="000000"/><w:bottom w:val="single" w:sz="8" w:space="0" w:color="000000"/><w:insideH w:val="single" w:sz="2" w:space="0" w:color="999999"/></w:tblBorders>')
        tblPr.append(borders)
        mar = parse_xml(f'<w:tblCellMar xmlns:w="{W}"><w:top w:w="15" w:type="dxa"/><w:left w:w="50" w:type="dxa"/><w:bottom w:w="15" w:type="dxa"/><w:right w:w="50" w:type="dxa"/></w:tblCellMar>')
        tblPr.append(mar)
        layout = OxmlElement("w:tblLayout"); layout.set(qn("w:type"), "fixed"); tblPr.append(layout)
        for ci, w in enumerate(widths):
            t.columns[ci].width = Inches(w)
            for row in t.rows: row.cells[ci].width = Inches(w)
        for e in tblPr.findall(qn("w:tblW")): tblPr.remove(e)
        tw = OxmlElement("w:tblW"); tw.set(qn("w:w"), str(int(sum(widths) * 1440))); tw.set(qn("w:type"), "dxa"); tblPr.append(tw)
        ind = OxmlElement("w:tblInd"); ind.set(qn("w:w"), "0"); ind.set(qn("w:type"), "dxa"); tblPr.append(ind)
        def fill(cell, text, bold=False, al=None):
            p = cell.paragraphs[0]; p.style = self.doc.styles["Table Text IEEE"]
            if al is not None: p.alignment = al
            self.rich(p, text, size=7.5, bold=bold)
        for ci, h in enumerate(header):
            fill(t.rows[0].cells[ci], h, bold=True, al=(WD_ALIGN_PARAGRAPH.RIGHT if (align and align[ci] == 'r') else None))
            tcPr = t.rows[0].cells[ci]._tc.get_or_add_tcPr()
            tcPr.append(parse_xml(f'<w:tcBorders xmlns:w="{W}"><w:bottom w:val="single" w:sz="6" w:space="0" w:color="000000"/></w:tcBorders>'))
        for ri, row in enumerate(rows):
            for ci, v in enumerate(row):
                al = None
                if align and align[ci] == "r": al = WD_ALIGN_PARAGRAPH.RIGHT
                fill(t.rows[ri + 1].cells[ci], v, al=al)
        for ri_, row in enumerate(t.rows):
            trPr = row._tr.get_or_add_trPr(); trPr.append(OxmlElement("w:cantSplit"))
            if ri_ < len(t.rows) - 1:
                for c_ in row.cells:
                    for p_ in c_.paragraphs: p_.paragraph_format.keep_with_next = True
        hdr = t.rows[0]._tr.get_or_add_trPr(); hdr.append(OxmlElement("w:tblHeader"))
        sp = self.doc.add_paragraph(); sp.paragraph_format.space_after = Pt(4)
        for r_ in sp.runs: r_.font.size = Pt(2)
        sp.paragraph_format.line_spacing = Pt(4)
        return self.tab_n

    # ---------- equations (OMML) ----------
    @staticmethod
    def _mr(text, plain=False):
        txt = text.replace("&", "&amp;").replace("<", "&lt;")
        sty = '<m:rPr><m:sty m:val="p"/></m:rPr>' if plain else ""
        return f'<m:r>{sty}<w:rPr><w:rFonts w:ascii="Cambria Math" w:hAnsi="Cambria Math"/></w:rPr><m:t xml:space="preserve">{txt}</m:t></m:r>'
    def eq(self, parts):
        """parts: list of ('t',text) | ('p',text) plain | ('sub',base,sub)"""
        self.eq_n += 1
        p = self.doc.add_paragraph(style="Equation IEEE")
        p.paragraph_format.tab_stops.add_tab_stop(Inches(1.75), WD_TAB_ALIGNMENT.CENTER)
        p.paragraph_format.tab_stops.add_tab_stop(Inches(3.5), WD_TAB_ALIGNMENT.RIGHT)
        p.add_run("\t")
        xml = ""
        for pt in parts:
            if pt[0] == "t": xml += self._mr(pt[1])
            elif pt[0] == "p": xml += self._mr(pt[1], plain=True)
            elif pt[0] == "sub": xml += f'<m:sSub><m:e>{self._mr(pt[1])}</m:e><m:sub>{self._mr(pt[2])}</m:sub></m:sSub>'
        om = parse_xml(f'<m:oMath xmlns:m="{M}" xmlns:w="{W}">{xml}</m:oMath>')
        p._p.append(om)
        r = p.add_run(f"\t({self.eq_n})"); r.font.name = FONT
        return self.eq_n

    # ---------- references ----------
    def references(self):
        self.h1("References")
        for i, k in enumerate(self.order):
            p = self.doc.add_paragraph(style="Reference IEEE")
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.tab_stops.add_tab_stop(Inches(0.3))
            p.add_run(f"[{i+1}]\t")
            self.rich(p, re.sub(r'"([^"]+)"', "\u201c\\1\u201d", self.refs[k]), size=8)
